"""Offline regression checks; no production credentials or services are used."""
import io
import tempfile
import unittest
from datetime import datetime, timezone, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from supabase_auth.errors import AuthApiError

from main import app
from api.auth import require_user
from api import routes
from services.data_processor import process_grant_rows, GrantDataValidationError
from services.sheets_reader import GoogleSheetsReadError
from services.word_generator import create_word_report, REPORT_FIELDS
from services.template_report import create_template_word_report, CHAPTER_TEMPLATES


class RegressionTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        self.grant = {'grant_number': 'DEMO-1', 'country': 'Canada',
                      'chapter': 'TIH UAE', 'supplier': 'Supplier', 'item': 'Pump',
                      'current_payment_usd': '125.50', 'no_of_beneficiaries': '10'}

    def tearDown(self):
        app.dependency_overrides.clear()

    def sign_in(self):
        app.dependency_overrides[require_user] = lambda: SimpleNamespace(id='test-user')

    def test_every_data_route_requires_authentication(self):
        for method, path in [('GET', '/api/grants'), ('GET', '/api/grants/test'),
                             ('POST', '/api/grants'), ('PUT', '/api/grants/test'),
                             ('POST', '/api/reports/pdf/test'), ('POST', '/api/reports/word/test'),
                             ('GET', '/api/reports/history')]:
            with self.subTest(path=path, method=method):
                self.assertEqual(self.client.request(method, path).status_code, 401)
        self.assertEqual(self.client.get('/health').status_code, 200)

    def test_invalid_token_is_rejected(self):
        with patch('api.auth._auth_client') as auth:
            auth.return_value.auth.get_user.side_effect = AuthApiError('Invalid token', 401, 'bad_jwt')
            self.assertEqual(self.client.get('/api/grants', headers={'Authorization': 'Bearer invalid'}).status_code, 401)

    def test_verified_token_can_read_grants(self):
        with patch('api.auth._auth_client') as auth, patch.object(routes, 'fetch_grant_rows_with_source', return_value=([], True)), patch.object(routes.supabase_service, 'cache_sheet_data'):
            auth.return_value.auth.get_user.return_value = SimpleNamespace(user=SimpleNamespace(id='user'))
            response = self.client.get('/api/grants', headers={'Authorization': 'Bearer verified'})
            self.assertEqual(response.status_code, 200)
            auth.return_value.auth.get_user.assert_called_once_with('verified')

    def test_invalid_numeric_cell_is_reported_without_hiding_other_grants(self):
        self.sign_in()
        invalid = dict(self.grant, current_payment_usd='12.3.4')
        valid = dict(self.grant, grant_number='DEMO-2')
        with patch.object(routes, 'fetch_grant_rows_with_source', return_value=([invalid, valid], True)), patch.object(routes.supabase_service, 'cache_sheet_data'):
            response = self.client.get('/api/grants')
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual([g['grant_number'] for g in body['grants']], ['DEMO-1', 'DEMO-2'])
        self.assertIsNone(body['grants'][0]['current_payment_usd'])
        self.assertEqual(len(body['row_errors']), 1)
        self.assertIn('Row 2 (DEMO-1)', body['row_errors'][0])

    def test_duplicate_grant_numbers_are_reported(self):
        summary = process_grant_rows([self.grant, dict(self.grant, supplier='Other')])
        self.assertEqual(len(summary.grants), 1)
        self.assertEqual(summary.grants[0].supplier, 'Supplier')
        self.assertIn('already appears in row 2', summary.row_errors[0])

    def test_blank_and_zero_numbers_do_not_count_as_filled(self):
        from services.data_processor import REQUIRED_FOR_STATUS, calculate_report_status
        full = {field: 'x' for field in REQUIRED_FOR_STATUS}
        full.update(total_grant_amount_orig='10', total_grant_amount_usd='10', quantity='1',
                    no_of_beneficiaries='5', remaining_payment_usd='0', remaining_payment_orig='0')
        self.assertEqual(calculate_report_status(full), 'Pending')
        self.assertEqual(calculate_report_status(dict(full, quantity='0')), 'Incomplete Information')
        self.assertEqual(calculate_report_status(dict(full, quantity=None)), 'Incomplete Information')
        self.assertEqual(calculate_report_status(dict(full, link_to_utilization_report='https://x')), 'Complete')
        self.assertIsNone(process_grant_rows([dict(self.grant, quantity='')]).grants[0].quantity)

    def test_exchange_rate_is_read_and_validated(self):
        from services.data_processor import validate_grant_dict
        self.assertEqual(process_grant_rows([dict(self.grant, exchange_rate='0.73')]).grants[0].exchange_rate, 0.73)
        self.assertIsNone(process_grant_rows([self.grant]).grants[0].exchange_rate)
        valid = dict(self.grant, chapter='IDF', project_type='Expansion', payment_status='Pending')
        validate_grant_dict(dict(valid, exchange_rate=0.73))
        with self.assertRaises(GrantDataValidationError):
            validate_grant_dict(dict(valid, exchange_rate=0))

    def test_report_history_saves_without_new_columns_before_migration(self):
        from postgrest.exceptions import APIError
        from services import supabase_service
        inserted = []
        def insert(payload):
            inserted.append(payload)
            request = unittest.mock.MagicMock()
            if 'grant_number' in payload:
                request.execute.side_effect = APIError({'code': 'PGRST204', 'message': "Could not find the 'grant_number' column"})
            else:
                request.execute.return_value.data = [payload]
            return request
        client = unittest.mock.MagicMock()
        client.table.return_value.insert.side_effect = insert
        with patch.object(supabase_service, 'get_supabase_client', return_value=client), self.assertLogs(level='WARNING'):
            saved = supabase_service.save_report_metadata(generated_by='u', report_format='PDF', row_count=1,
                                                          total_received=1, total_beneficiaries=1, grant_number='DEMO-1')
        self.assertEqual(len(inserted), 2)
        self.assertNotIn('grant_number', saved)

    def test_version_changes_when_the_row_changes(self):
        first = process_grant_rows([self.grant]).grants[0].version
        self.assertEqual(first, process_grant_rows([dict(self.grant)]).grants[0].version)
        self.assertNotEqual(first, process_grant_rows([dict(self.grant, item='Valve')]).grants[0].version)

    def test_recent_fallback_is_labeled_and_expired_cache_rejected(self):
        self.sign_in()
        for hours, expected in [(1, 200), (48, 503)]:
            cached = {'cached_at': (datetime.now(timezone.utc) - timedelta(hours=hours)).isoformat(), 'data': [self.grant]}
            with patch.object(routes, 'fetch_grant_rows_with_source', side_effect=GoogleSheetsReadError('offline')), patch.object(routes.supabase_service, 'get_latest_sheet_cache', return_value=cached):
                response = self.client.get('/api/grants')
            self.assertEqual(response.status_code, expected)
            if expected == 200:
                self.assertIn('cache_warning', response.json())

    def test_generic_word_sections_and_current_schema(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'report.docx'
            sections = {key: key == 'financial' for key in REPORT_FIELDS}
            create_word_report(self.grant, path, sections)
            document = Document(path)
            text = '\n'.join(p.text for p in document.paragraphs)
            self.assertIn('Financial Summary', text)
            self.assertNotIn('Shipping & Documents', text)
            cells = ' '.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn('125.50', cells)

    def test_currency_formats_and_invalid_numbers(self):
        for value, expected in [('USD 1,250.50', 1250.50), ('1e3', 1000), ('0', 0)]:
            summary = process_grant_rows([dict(self.grant, current_payment_usd=value)])
            self.assertEqual(summary.grants[0].current_payment_usd, expected)
        for value in ['NaN', 'inf', 'USD nonsense', '1e309']:
            summary = process_grant_rows([dict(self.grant, current_payment_usd=value)])
            self.assertIsNone(summary.grants[0].current_payment_usd)
            self.assertEqual(len(summary.row_errors), 1)

    def test_all_chapter_templates_generate_valid_word_documents(self):
        with tempfile.TemporaryDirectory() as directory:
            for chapter, template in CHAPTER_TEMPLATES.items():
                for project_type in template if isinstance(template, dict) else ['Expansion']:
                    with self.subTest(chapter=chapter, project_type=project_type):
                        path = Path(directory) / 'template.docx'
                        create_template_word_report(dict(self.grant, chapter=chapter, project_type=project_type, supplier='Supplier & <Partners>'), path)
                        self.assertGreater(len(Document(path).paragraphs), 0)

    def test_cors_rejects_unlisted_origin(self):
        response = self.client.options('/api/grants', headers={
            'Origin': 'https://unlisted.example', 'Access-Control-Request-Method': 'GET'})
        self.assertEqual(response.status_code, 400)

    def test_downloads_record_history_and_cleanup_files(self):
        self.sign_in()
        directories = []
        original = tempfile.mkdtemp
        def track_directory(*args, **kwargs):
            directory = original(*args, **kwargs)
            directories.append(Path(directory))
            return directory
        for format in ['word', 'pdf']:
            with patch.object(routes, '_get_single_grant', return_value=self.grant), patch.object(routes.supabase_service, 'save_report_metadata') as history, patch.object(routes.tempfile, 'mkdtemp', side_effect=track_directory):
                response = self.client.post(f'/api/reports/{format}/DEMO-1', json={'sections': {'shipping': False}})
            self.assertEqual(response.status_code, 200)
            history.assert_called_once()
            self.assertEqual(history.call_args.kwargs['grant_number'], 'DEMO-1')
            self.assertFalse(directories[-1].exists())
            if format == 'word':
                document = Document(io.BytesIO(response.content))
                self.assertNotIn('Shipping & Documents', '\n'.join(p.text for p in document.paragraphs))
            else:
                self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertNotEqual(directories[0], directories[1])

    def test_custom_sections_use_generic_layout_even_for_template_chapter(self):
        self.sign_in()
        with patch.object(routes, '_get_single_grant', return_value=dict(self.grant, chapter='IDF')), patch.object(routes.supabase_service, 'save_report_metadata'), patch.object(routes, 'create_template_word_report') as template:
            response = self.client.post('/api/reports/word/DEMO-1', json={'sections': {'financial': False}})
        self.assertEqual(response.status_code, 200)
        template.assert_not_called()

    def test_history_is_available_to_authenticated_users(self):
        self.sign_in()
        with patch.object(routes.supabase_service, 'get_report_history', return_value=[{'id': 'report'}]):
            self.assertEqual(self.client.get('/api/reports/history').json(), [{'id': 'report'}])

    def test_failed_report_cleans_temporary_files(self):
        self.sign_in()
        with tempfile.TemporaryDirectory() as parent:
            directory = Path(parent) / 'download'
            directory.mkdir()
            with patch.object(routes, '_get_single_grant', return_value=self.grant), patch.object(routes.tempfile, 'mkdtemp', return_value=str(directory)), patch.object(routes, 'create_word_report', side_effect=RuntimeError('failure')), self.assertLogs(level='ERROR'):
                response = self.client.post('/api/reports/word/DEMO-1')
            self.assertEqual(response.status_code, 503)
            self.assertFalse(directory.exists())


if __name__ == '__main__':
    unittest.main()
